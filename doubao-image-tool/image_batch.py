"""Collector image downloading, source-position mapping and persistent exports."""
import copy
import io
import json
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from time import monotonic
from urllib.error import HTTPError
from urllib.parse import urljoin, urlsplit
from urllib.request import HTTPRedirectHandler, Request, build_opener

from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from storage import save_result, validate_inputs

MAX_DOWNLOAD_BYTES = 16 * 1024 * 1024
# Inactivity/inner-loop limits are secondary; the owning process enforces the hard bound.
DOWNLOAD_TIMEOUT = 4
DOWNLOAD_ATTEMPT_SECONDS = 8
DOWNLOAD_PROCESS_SECONDS = 18
HEADERS = ['平台', '商品ID', '商品标题', '图片类型', 'SKU', '顺序', '原图URL', '本地转换路径', '状态', '原因']
FIELDS = ['platform', 'product_id', 'title', 'kind', 'sku', 'order', 'url', 'output_path', 'status', 'reason']


def validate_image_url(url):
    if not isinstance(url, str) or len(url) > 16384 or any(c.isspace() for c in url):
        raise ValueError('图片 URL 无效')
    parsed = urlsplit(url)
    host = (parsed.hostname or '').lower()
    if (parsed.scheme != 'https' or parsed.username is not None or parsed.password is not None
            or parsed.port not in (None, 443) or parsed.fragment
            or not any(host == domain or host.endswith('.' + domain) for domain in ('pddpic.com', 'alicdn.com'))):
        raise ValueError('仅支持 HTTPS 拼多多/阿里图片域名')
    return url


class NoRedirect(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


def _download_launcher():
    return ([sys.executable] if getattr(sys, 'frozen', False)
            else [sys.executable, str(Path(__file__).with_name('run.py'))])


def download_image(url):
    """Kill and reap the short-lived transport process if any network phase stalls."""
    validate_image_url(url)
    with tempfile.TemporaryDirectory(prefix='doubao-image-download-') as folder:
        output = Path(folder) / 'image.bin'
        try:
            result = subprocess.run(_download_launcher() + ['--download-image'],
                input=json.dumps({'url': url, 'output_path': str(output)}),
                capture_output=True, text=True, encoding='utf-8', check=False,
                timeout=DOWNLOAD_PROCESS_SECONDS,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        except subprocess.TimeoutExpired:
            # subprocess.run kills and waits for its direct child before raising.
            raise TimeoutError('图片下载超过总时限，请重试') from None
        try:
            metadata = json.loads(result.stdout)
            if result.returncode or metadata.get('extension') not in ('.png', '.jpg', '.webp'):
                if metadata.get('error') == 'invalid': raise ValueError('图片 URL 或图片内容无效')
                raise OSError('图片下载失败，请重试')
            if not output.is_file() or output.stat().st_size > MAX_DOWNLOAD_BYTES:
                raise ValueError('图片下载结果大小无效')
            data = output.read_bytes(); extension = metadata['extension']
            validate_inputs([('download' + extension, data)])
            return data, extension
        except (json.JSONDecodeError, KeyError, TypeError):
            raise OSError('图片下载进程未返回有效结果') from None


def download_helper_main():
    """Hidden source/frozen entry point; signed URLs are received only on stdin."""
    try:
        payload = json.loads(sys.stdin.read())
        data, extension = _download_image_transport(payload['url'])
        with Path(payload['output_path']).open('xb') as handle: handle.write(data)
        print(json.dumps({'extension': extension}), flush=True)
        return 0
    except Exception as exc:
        # Avoid echoing signed URLs, query parameters or filesystem paths from errors.
        print(json.dumps({'error': 'invalid' if isinstance(exc, ValueError) else 'download'}), flush=True)
        return 1


def _download_image_transport(url):
    """Validate every redirect; the parent process interrupts stalled headers or reads."""
    validate_image_url(url)
    opener = build_opener(NoRedirect())
    for attempt in range(2):
        current = url
        deadline = monotonic() + DOWNLOAD_ATTEMPT_SECONDS
        try:
            for hop in range(6):
                if monotonic() >= deadline: raise TimeoutError('图片下载超时')
                try:
                    response = opener.open(Request(current, headers={'User-Agent': 'Mozilla/5.0',
                        'Accept': 'image/png,image/jpeg,image/webp'}), timeout=DOWNLOAD_TIMEOUT)
                except HTTPError as exc:
                    if exc.code not in (301, 302, 303, 307, 308):
                        raise
                    location = exc.headers.get('Location')
                    exc.close()
                    if not location or hop == 5: raise ValueError('图片重定向次数超限或缺少目标')
                    current = validate_image_url(urljoin(current, location))
                    continue
                with response:
                    validate_image_url(response.geturl())
                    if response.headers.get('Content-Length') and int(response.headers['Content-Length']) > MAX_DOWNLOAD_BYTES:
                        raise ValueError('图片大小超过下载限制')
                    chunks = []; size = 0
                    while True:
                        if monotonic() >= deadline: raise TimeoutError('图片下载超时')
                        chunk = response.read1(min(65536, MAX_DOWNLOAD_BYTES + 1 - size))
                        if not chunk: break
                        chunks.append(chunk); size += len(chunk)
                        if size > MAX_DOWNLOAD_BYTES: raise ValueError('图片大小超过下载限制')
                    data = b''.join(chunks)
                try:
                    with Image.open(io.BytesIO(data)) as image:
                        extension = {'PNG': '.png', 'JPEG': '.jpg', 'WEBP': '.webp'}.get(image.format)
                    if not extension: raise ValueError('不支持图片格式')
                    validate_inputs([('download' + extension, data)])
                except Exception as exc:
                    raise ValueError('下载内容不是有效静态图片或尺寸超过限制') from exc
                return data, extension
        except ValueError:
            raise
        except (OSError, HTTPError):
            if attempt == 1: raise


def build_items(entries):
    if not isinstance(entries, list) or not entries: raise ValueError('请选择商品图片')
    items = []; by_url = {}
    for position, source in enumerate(entries):
        if not isinstance(source, dict): raise ValueError('图片位置格式无效')
        ref = copy.deepcopy(source)
        for key in ('platform', 'product_id', 'title', 'sku', 'url', 'product_url'):
            if not isinstance(ref.get(key, ''), str): raise ValueError('图片元数据必须是文本')
            ref.setdefault(key, '')
        if not ref['platform'] or not ref['product_id'] or not ref['url']:
            raise ValueError('缺少平台、商品ID或图片URL')
        if ref.get('kind') not in ('main', 'detail', 'sku') or type(ref.get('order')) is not int or ref['order'] < 1:
            raise ValueError('图片类型或顺序无效')
        ref.update(position=position, output_path=None)
        item = by_url.get(ref['url'])
        if item is None:
            item = {'index': len(items), 'name': f'image_{len(items) + 1}.png', 'url': ref['url'],
                    'refs': [], 'status': 'queued', 'phase': 'downloading', 'message': '',
                    'result': None, 'review_needed': True}
            items.append(item); by_url[ref['url']] = item
        item['refs'].append(ref)
    return items


def counts(job):
    unique = [item for item in job['items'] if 'alias_of' not in item]
    return {'total': sum(len(item['refs']) for item in job['items']), 'unique': len(unique),
            'downloaded': sum(bool(item.get('sha256')) for item in unique),
            'converted': sum(item['status'] == 'completed' for item in unique),
            'failed': sum(item['status'] == 'failed' for item in unique)}


def safe_component(value, limit=70):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(value)).strip(' .')[:limit].strip(' .') or 'image'
    if re.fullmatch(r'(?i)(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?', value): value = '_' + value
    return value


def save_positions(job, refs, data):
    """Use exclusive result writes; one destination for every original source position."""
    results = []
    for ref in refs:
        folder = Path(job['run_dir']) / (safe_component(ref['platform'], 15) + '_' +
            safe_component(ref['product_id'], 35) + '_' + safe_component(ref['title'], 45))
        name = f"{ref['kind']}_{ref['order']:03d}"
        if ref['sku']: name += '_' + safe_component(ref['sku'], 55)
        name += f"_pos{ref['position'] + 1:04d}.png"
        result = save_result(folder, name, data, job['prompt'])
        ref['output_path'] = result['output_path']
        results.append(result)
    return results[0]


def write_report(job):
    """Every source row remains present, including queued, paused and failed images."""
    rows = []
    for item in job['items']:
        for ref in item['refs']:
            row = copy.deepcopy(ref)
            row.update(status=item['status'], reason=item['message'])
            rows.append(row)
    rows.sort(key=lambda row: row['position'])
    mapping = Path(job['mapping_path']); temp = mapping.with_suffix('.json.tmp')
    temp.write_text(json.dumps({'job_id': job['id'], 'source_task_id': job['source_task_id'],
        'status': job['status'], 'counts': counts(job), 'rows': rows}, ensure_ascii=False, indent=2), encoding='utf-8')
    temp.replace(mapping)
    book = Workbook(); sheet = book.active; sheet.title = '图片转换清单'
    sheet.append(HEADERS)
    for row in rows:
        sheet.append([row.get(key) for key in FIELDS])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str): cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.row_dimensions[sheet.max_row].height = 42
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='24476A')
    for column, width in zip('ABCDEFGHIJ', (14, 22, 40, 14, 28, 10, 55, 65, 18, 42)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
    path = Path(job['manifest_path']); pending = path.with_suffix('.xlsx.tmp')
    book.save(pending); pending.replace(path)
